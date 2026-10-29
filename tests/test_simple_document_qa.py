"""Conversational questions, long-source extraction, and grounded concise answers."""

import copy
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from medical_doc_assistant.grounded_qa import DocumentAssistant, MISSING_ANSWER
from medical_doc_assistant.ingestion import read_document_file, source_sections
from medical_doc_assistant.qa_language import normalize_question, question_intent
from medical_doc_assistant.rag_evaluation import response_supported
from test_grounded_qa import FakeAnswerability, FakeEmbedder, FakeGenerator

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures" / "qa"
CONFIG = json.loads((ROOT / "configs" / "phase4.json").read_text(encoding="utf-8"))
SETTINGS = json.loads((ROOT / "configs" / "phase5.json").read_text(encoding="utf-8"))


class SimpleQuestionTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix=".test-", dir=ROOT)
        self.root = Path(self.temporary.name)
        self.config = copy.deepcopy(CONFIG)
        self.config["rag"].update(chunk_tokens=80, overlap_tokens=15)
        self.assistant = DocumentAssistant(self.root, self.config, embedder=FakeEmbedder(self.root),
                                           answerability=FakeAnswerability(), generator=FakeGenerator("Yes, it is always safe. Take 999 mg."))
        self.text = (FIXTURES / "long_patient_leaflet.txt").read_text(encoding="utf-8")
        self.index = self.assistant.index_document(self.text, "fictional Lumenol test leaflet")

    def tearDown(self):
        self.temporary.cleanup()

    def answered(self, question, index=None):
        index = index or self.index
        response = self.assistant.answer(index, question)
        self.assertEqual(response["status"], "answered", response)
        self.assertTrue(response_supported(response, index.document), response)
        self.assertNotIn("999", response["answer"])
        return response

    def missing(self, question, index=None):
        response = self.assistant.answer(index or self.index, question)
        self.assertEqual(response["status"], "not_found", response)
        self.assertEqual(response["answer"], MISSING_ANSWER)
        self.assertEqual(response["citations"], [])
        return response

    def test_fixture_has_many_chunks_and_relevant_information_after_old_excerpt_limit(self):
        self.assertGreater(len(self.text.split()), 1500)
        self.assertGreater(len(self.index.chunks), 25)
        sections = {section["field"]: section["text"] for section in source_sections(self.index.document)}
        self.assertGreater(sections["pregnancy"].index("Lumenol should be used"), 1800)
        self.assertGreater(sections["dosage_and_administration"].index("For adults 18"), 1800)
        self.answered("pregnancy")
        self.answered("dosage")

    def test_contractions_and_casual_pregnancy_questions_are_equivalent(self):
        questions = ("can I take the medicine if I have pregnancy ?", "can I take the medicine if I'm pregnant ?",
                     "Can I take it if I’m pregnant?", "I'm pregnant, can I take this medicine?",
                     "can I take it if im pregnant", "Is it safe while I am pregnant?",
                     "Is it okay if I'm pregnant?", "Can I use this if I am expecting a baby?",
                     "pregnancy", "pregnancy?", "Tell me about pregnancy please")
        for question in questions:
            with self.subTest(question=question):
                response = self.answered(question)
                self.assertIn("only if the potential benefit justifies the potential risk to the fetus", response["answer"])
                self.assertIn("no adequate controlled human studies", response["answer"].lower())
                self.assertLess(len(response["answer"]), 500)
                self.assertFalse(response["answer"].lower().startswith("yes"))

    def test_one_word_dosage_and_amount_questions_do_not_return_packaging_strengths(self):
        for question in ("dosage", "dose", "directions", "Dosage?", "how much should I take?", "How many pills do I take?"):
            with self.subTest(question=question):
                response = self.answered(question)
                self.assertIn("one 5 mg tablet", response["answer"])
                self.assertIn("18 years and older", response["answer"])
                self.assertIn("do not exceed 10 mg in 24 hours", response["answer"])
                self.assertNotIn("100 mg", response["answer"])
                self.assertNotIn("Printed directions can be separated", response["answer"])
                self.assertLess(len(response["answer"]), 500)

    def test_frequency_and_course_duration_are_distinct(self):
        response = self.answered("how often should I take it?")
        self.assertIn("twice daily", response["answer"])
        self.assertNotIn("7 days", response["answer"])
        response = self.answered("how long should I take it?")
        self.assertIn("for 7 days only", response["answer"])
        self.assertIn("unless the prescriber gives different instructions", response["answer"])
        self.assertNotIn("twice daily", response["answer"])

    def test_maximum_dose_keeps_its_age_and_quantity_conditions(self):
        response = self.answered("maximum dosage")
        self.assertIn("18 years and older", response["answer"])
        self.assertIn("do not exceed 10 mg in 24 hours", response["answer"])

    def test_child_question_does_not_substitute_adult_directions(self):
        for question in ("dosage for children", "how much should a child take?", "dose for kids"):
            with self.subTest(question=question):
                response = self.answered(question)
                self.assertIn("dosing has not been established", response["answer"])
                self.assertIn("do not use the adult dosage", response["answer"])
                self.assertNotIn("one 5 mg tablet", response["answer"])

    def test_food_permission_and_crushing_prohibition_have_grounded_yes_no(self):
        response = self.answered("Can I take this with food?")
        self.assertEqual(response["answer_polarity"], "yes")
        self.assertTrue(response["answer"].startswith("Yes, according to the document"))
        self.assertIn("with food", response["answer"])
        response = self.answered("Can I crush the tablet?")
        self.assertEqual(response["answer_polarity"], "no")
        self.assertIn("Do not crush or chew", response["answer"])

    def test_breastfeeding_uses_nursing_mothers_not_pregnancy_conditions(self):
        response = self.answered("Can I breastfeed while taking the medicine?")
        self.assertEqual(response["answer_polarity"], "no")
        self.assertIn("Nursing should not be undertaken", response["answer"])
        self.assertNotIn("potential benefit", response["answer"])

    def test_other_short_section_topics_return_brief_relevant_statements(self):
        for question, expected in (("side effects", "dizziness, nausea, and fatigue"),
                                   ("warnings", "Do not use Lumenol if you are allergic"),
                                   ("what is this medicine used for?", "amber-indicator symptoms"),
                                   ("storage", "20 to 25 degrees Celsius")):
            with self.subTest(question=question):
                response = self.answered(question)
                self.assertIn(expected, response["answer"])
                self.assertLess(len(response["answer"]), 600)

    def test_price_answers_when_present_and_falls_back_when_absent(self):
        for question in ("price", "how much does it cost?", "What is the retail price stated in the document?"):
            with self.subTest(question=question):
                self.assertIn("$24.95", self.answered(question)["answer"])
        other = self.assistant.index_document((FIXTURES / "long_uncertain_leaflet.txt").read_text(encoding="utf-8"))
        response = self.missing("price", other)
        self.assertIn("No retail price", response["missing_reason"])

    def test_personal_missing_facts_do_not_return_dates_or_numbers_from_other_sections(self):
        for question in ("What is my insurance policy number?", "What is my prescribing doctor's telephone number?", "expiration date"):
            with self.subTest(question=question):
                self.missing(question)

    def test_absent_drug_is_not_supported_by_an_unrelated_mention_in_description(self):
        for question in ("Can I take amoxicillin if I'm pregnant?", "amoxicillin dosage"):
            with self.subTest(question=question):
                response = self.missing(question)
                self.assertIn("specific medicine", response["missing_reason"])

    def test_animal_findings_do_not_turn_uncertain_human_evidence_into_yes(self):
        text = (FIXTURES / "long_uncertain_leaflet.txt").read_text(encoding="utf-8")
        index = self.assistant.index_document(text)
        response = self.answered("Can I take it if I'm pregnant?", index)
        self.assertEqual(response["answer_polarity"], "uncertain")
        self.assertIn("Safety during human pregnancy has not been established", response["answer"])
        self.assertFalse(response["answer"].startswith("Yes"))
        self.assertNotIn("only if the potential benefit", response["answer"])
        self.missing("dosage", index)

    def test_explicit_pregnancy_prohibition_and_permission_are_distinguished(self):
        for text, polarity in (("Pregnancy\nDo not use this fictional product during pregnancy.", "no"),
                               ("Pregnancy\nThis fictional source explicitly states that the product may be used during pregnancy.", "yes")):
            with self.subTest(polarity=polarity):
                index = self.assistant.index_document(text)
                response = self.answered("Can I take it if I am pregnant?", index)
                self.assertEqual(response["answer_polarity"], polarity)

    def test_negated_question_preserves_source_without_ambiguous_yes_no_prefix(self):
        index = self.assistant.index_document("Pregnancy\nDo not use this fictional product during pregnancy.")
        response = self.answered("Shouldn't I take it while pregnant?", index)
        self.assertIsNone(response["answer_polarity"])
        self.assertEqual(response["answer"], "Do not use this fictional product during pregnancy.")
        self.assertIn("not", normalize_question("Shouldn't I take amoxicillin?"))

    def test_absolute_prohibition_is_not_softened_by_another_consult_doctor_sentence(self):
        index = self.assistant.index_document("Pregnancy\nDo not use this fictional product during pregnancy.\nAsk your doctor about pregnancy before use.")
        response = self.answered("Can I take it if I'm pregnant?", index)
        self.assertEqual(response["answer_polarity"], "no")
        self.assertIn("Do not use", response["answer"])

    def test_unless_condition_is_not_presented_as_an_unconditional_no(self):
        index = self.assistant.index_document("Pregnancy\nDo not use during pregnancy unless the prescriber confirms that benefit outweighs risk.")
        response = self.answered("Can I take it if I'm pregnant?", index)
        self.assertEqual(response["answer_polarity"], "conditional")
        self.assertIn("unless the prescriber confirms", response["answer"])

    def test_pregnancy_specific_dose_is_not_invented_from_general_directions(self):
        self.missing("What dosage should I take if I'm pregnant?")

    def test_standard_dose_abbreviations_are_simplified_but_evidence_is_original(self):
        index = self.assistant.index_document("Dosage\nThe recommended initial dose is 5 mg b.i.d. with food.")
        response = self.answered("dosage", index)
        self.assertIn("twice a day", response["answer"])
        self.assertIn("b.i.d.", response["citations"][0]["quote"])

    def test_full_section_is_available_only_when_explicitly_requested(self):
        short = self.answered("dosage")
        full = self.answered("show the full dosage section")
        self.assertEqual(full["answer_style"], "full_section")
        self.assertGreater(len(full["answer"]), len(short["answer"]) * 5)
        self.assertIn("Printed directions can be separated", full["answer"])

    def test_docx_numbered_headings_and_aliases_keep_quote_offsets(self):
        document = Document()
        document.add_heading("2.1 Dosage:", level=1)
        document.add_paragraph("Adults 18 years and older:")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Take one 2.5 mg tablet twice daily"
        table.cell(0, 1).text = "only if the prescriber advises; do not exceed 5 mg per day."
        document.add_heading("8.1 Pregnancy and breast-feeding:", level=1)
        document.add_paragraph("Do not use this fictional product during pregnancy.")
        buffer = io.BytesIO()
        document.save(buffer)
        text = read_document_file("test.docx", buffer.getvalue(), SETTINGS)
        index = self.assistant.index_document(text)
        response = self.answered("dosage", index)
        self.assertIn("18 years and older", response["answer"])
        self.assertIn("2.5 mg", response["answer"])
        self.assertIn("only if", response["answer"])
        self.assertEqual(self.answered("Can I take it if I'm pregnant?", index)["answer_polarity"], "no")

    def test_present_values_for_personal_fields_are_supported_when_supplied(self):
        index = self.assistant.index_document("Patient record\nInsurance policy number: AB1234567.\nPrescriber's telephone: +1 202-555-0123.\nExpiration date: 2028-04-30.")
        for question, value in (("What is my insurance policy number?", "AB1234567"),
                                ("What is my prescribing doctor's telephone number?", "+1 202-555-0123"),
                                ("expiration date", "2028-04-30")):
            with self.subTest(question=question):
                self.assertIn(value, self.answered(question, index)["answer"])

    def test_answer_verifier_rejects_tampered_polarity_and_unsupported_summary(self):
        response = self.answered("Can I crush the tablet?")
        altered = copy.deepcopy(response)
        altered["answer_polarity"] = "yes"
        altered["answer"] = "Yes, it is safe."
        self.assertFalse(response_supported(altered, self.index.document))
        altered = copy.deepcopy(response)
        altered["answer"] += " You can take 999 mg."
        self.assertFalse(response_supported(altered, self.index.document))

    def test_question_intent_does_not_confuse_onset_with_treatment_duration(self):
        self.assertIsNone(question_intent("How long can this product take for full effect?").topic)
        self.assertEqual(question_intent("how long should I take it?").detail, "duration")


if __name__ == "__main__":
    unittest.main()
