"""Follow-up memory, overview, dynamic suggestions, and source navigation."""

import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from reportlab.pdfgen.canvas import Canvas

from medical_doc_assistant.api import create_app
from medical_doc_assistant.grounded_qa import DocumentAssistant
from medical_doc_assistant.rag_evaluation import response_supported
from medical_doc_assistant.serving import ApplicationRuntime
from medical_doc_assistant.source_navigation import highlighted_text
from test_grounded_qa import FakeAnswerability, FakeEmbedder, FakeGenerator

ROOT = Path(__file__).resolve().parents[1]
CONFIG = json.loads((ROOT / "configs" / "phase4.json").read_text(encoding="utf-8"))
TEXT = "Medicine name: Testolin\nActive ingredient: test compound 5 mg\nDescription\nEach tablet contains test compound 5 mg.\nDosage\nAdults: take one 5 mg tablet twice daily with food; do not exceed 10 mg per day.\nChildren: dosing has not been established; do not use adult instructions.\nPregnancy\nUse during pregnancy only if the prescriber confirms that benefit outweighs risk.\nWarnings\nDo not use if allergic to the listed ingredient.\nSide effects\nReported reactions include nausea and fatigue."
CLASSIFICATION = {"status": "classified", "predicted_class": "acid_reducer", "model": "fixture",
                  "model_name": "Fixture classifier", "latency_ms": 0}


class DocumentExperienceTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix=".test-", dir=ROOT)
        self.runtime = ApplicationRuntime(ROOT, Path(self.directory.name))
        self.runtime.assistant = DocumentAssistant(ROOT, CONFIG, embedder=FakeEmbedder(Path(self.directory.name)),
            answerability=FakeAnswerability(), generator=FakeGenerator("unsupported"), storage_root=Path(self.directory.name))
        self.classifier = patch.object(self.runtime, "classify", return_value=copy.deepcopy(CLASSIFICATION))
        self.classifier.start()
        self.client = TestClient(create_app(self.runtime))

    def tearDown(self):
        self.client.close()
        self.classifier.stop()
        self.directory.cleanup()

    def upload_text(self, text=TEXT):
        response = self.client.post("/documents", json={"text": text, "source_name": "Testolin.txt"})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def ask(self, document, question):
        response = self.client.post("/ask", json={"document_id": document["document_id"], "question": question})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_follow_up_frequency_and_children_reuse_the_dosage_topic(self):
        document = self.upload_text()
        self.ask(document, "dosage")
        frequency = self.ask(document, "how often?")
        self.assertTrue(frequency["context_used"])
        self.assertEqual(frequency["context_topic"], "dosage")
        self.assertIn("twice daily", frequency["answer"])
        children = self.ask(document, "what about children?")
        self.assertTrue(children["context_used"])
        self.assertIn("not been established", children["answer"])
        self.assertNotIn("one 5 mg tablet", children["answer"])
        self.assertTrue(response_supported(children, document["text"]))

    def test_follow_up_safety_reuses_pregnancy_but_not_another_document(self):
        first = self.upload_text()
        self.ask(first, "pregnancy")
        answer = self.ask(first, "is it safe?")
        self.assertTrue(answer["context_used"])
        self.assertEqual(answer["answer_polarity"], "conditional")
        second = self.upload_text("Dosage\nTake one 2.5 mg tablet daily.\nWarnings\nKeep the leaflet.")
        clarification = self.ask(second, "is it safe?")
        self.assertEqual(clarification["status"], "clarification_needed")
        self.assertFalse(clarification["context_used"])

    def test_overview_is_source_backed_and_detects_the_edited_name_mismatch(self):
        document = self.upload_text()
        self.assertEqual(document["overview"]["identity"]["medicine_name"]["text"], "Testolin")
        self.assertIn("test compound", document["overview"]["identity"]["active_ingredient"]["text"])
        for field in [*document["overview"]["identity"].values(), *document["overview"]["summary"].values()]:
            for citation in field["citations"]:
                self.assertEqual(document["text"][citation["char_start"]:citation["char_end"]], citation["quote"])
                self.assertTrue(citation["locations"])
        path = ROOT / "upload_examples" / "docx" / "05_Desoximetasone.docx"
        edited = self.client.post("/documents/upload", files={"file": (path.name, path.read_bytes())}).json()
        self.assertTrue(any("pindolol" in issue["message"].lower() for issue in edited["overview"]["consistency_issues"]))

    def test_suggestions_only_include_answerable_document_topics(self):
        document = self.upload_text()
        labels = [item["label"] for item in document["suggestions"]]
        self.assertIn("Dosage", labels)
        self.assertIn("Pregnancy", labels)
        self.assertNotIn("Price", labels)
        self.assertNotIn("Breastfeeding", labels)
        for suggestion in document["suggestions"]:
            self.assertEqual(self.ask(document, suggestion["question"])["status"], "answered")

    def test_docx_citations_have_paragraph_or_table_locations(self):
        source = Document()
        source.add_heading("Dosage", level=1)
        table = source.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Adults"
        table.cell(0, 1).text = "Take one 2.5 mg tablet twice daily."
        source.add_heading("Pregnancy", level=1)
        source.add_paragraph("Do not use during pregnancy.")
        buffer = io.BytesIO()
        source.save(buffer)
        document = self.client.post("/documents/upload", files={"file": ("source.docx", buffer.getvalue())}).json()
        dosage = self.ask(document, "dosage")
        self.assertTrue(any(location.get("table") == 1 for citation in dosage["citations"] for location in citation["locations"]))
        pregnancy = self.ask(document, "pregnancy")
        self.assertTrue(any(location.get("paragraph") == 3 for citation in pregnancy["citations"] for location in citation["locations"]))

    def test_pdf_citation_has_page_and_original_can_be_opened(self):
        buffer = io.BytesIO()
        canvas = Canvas(buffer)
        canvas.drawString(50, 750, "Description")
        canvas.drawString(50, 725, "A fictional tablet for testing.")
        canvas.showPage()
        canvas.drawString(50, 750, "Dosage")
        canvas.drawString(50, 725, "Adults: take one 5 mg tablet twice daily.")
        canvas.save()
        document = self.client.post("/documents/upload", files={"file": ("source.pdf", buffer.getvalue())}).json()
        answer = self.ask(document, "dosage")
        self.assertTrue(any(location.get("page") == 2 for citation in answer["citations"] for location in citation["locations"]))
        original = self.client.get("/documents/" + document["document_id"] + "/original")
        self.assertEqual(original.content, buffer.getvalue())

    def test_highlighted_source_escapes_markup_and_marks_exact_span(self):
        rendered = highlighted_text("<script>bad()</script> Dose 2.5 mg", highlights=[(22, 33)])
        self.assertNotIn("<script>", rendered)
        self.assertIn("&lt;script&gt;", rendered)
        self.assertIn("<mark>", rendered)
        with self.assertRaises(ValueError):
            highlighted_text("short", start=20)


if __name__ == "__main__":
    unittest.main()
