"""Capstone ingestion, source-offset, document-isolation, and API integration tests."""

import copy
import io
import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from medical_doc_assistant.api import create_app
from medical_doc_assistant.ingestion import extract_source_fields, read_document_file
from medical_doc_assistant.serving import ApplicationRuntime
from medical_doc_assistant.grounded_qa import DocumentAssistant
from medical_doc_assistant.rag_evaluation import response_supported
from test_grounded_qa import FakeAnswerability, FakeEmbedder, FakeGenerator

ROOT = Path(__file__).resolve().parents[1]
SETTINGS = json.loads((ROOT / "configs" / "phase5.json").read_text(encoding="utf-8"))
TEXT = "Dosage\nUse 2.5 units only if age < 12. Do NOT exceed the stated limit.\nWarnings\nKeep the amber widget dry.\nSide Effects\nThe source reports a blue indicator."


class IngestionTests(unittest.TestCase):
    def test_unicode_numbers_negation_and_offsets_survive_text_ingestion(self):
        text = read_document_file("source.txt", ("\ufeff" + TEXT).encode("utf-8"), SETTINGS)
        fields = extract_source_fields(text)
        self.assertIn("2.5", fields["fields"]["dosage"]["text"])
        self.assertIn("Do NOT", fields["fields"]["dosage"]["text"])
        self.assertIn("< 12", fields["fields"]["dosage"]["text"])
        for section in fields["sections"]:
            self.assertEqual(text[section["char_start"]:section["char_end"]], section["text"])

    def test_missing_fact_sections_are_not_invented(self):
        fields = extract_source_fields("Warnings: Keep the amber widget dry.")
        self.assertTrue(fields["fields"]["warnings"]["present"])
        self.assertFalse(fields["fields"]["dosage"]["present"])
        self.assertEqual(fields["fields"]["side_effects"]["text"], "")

    def test_docx_table_order_preserves_section_assignment(self):
        document = Document()
        document.add_heading("Dosage", level=1)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Example", "2.5 units"
        document.add_heading("Warnings", level=1)
        document.add_paragraph("Keep the amber widget dry.")
        buffer = io.BytesIO()
        document.save(buffer)
        text = read_document_file("source.docx", buffer.getvalue(), SETTINGS)
        fields = extract_source_fields(text)
        self.assertIn("2.5 units", fields["fields"]["dosage"]["text"])
        self.assertNotIn("2.5 units", fields["fields"]["warnings"]["text"])

    def test_image_only_pdf_returns_a_clear_text_extraction_error(self):
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        buffer = io.BytesIO()
        writer.write(buffer)
        with self.assertRaisesRegex(ValueError, "no extractable text"):
            read_document_file("scan.pdf", buffer.getvalue(), SETTINGS)

    def test_unsupported_and_oversized_files_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "text-based PDF"):
            read_document_file("image.exe", b"example", SETTINGS)
        settings = copy.deepcopy(SETTINGS)
        settings["maximum_upload_bytes"] = 4
        with self.assertRaisesRegex(ValueError, "limit"):
            read_document_file("source.txt", b"too long", settings)


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix=".test-", dir=ROOT)
        self.runtime = ApplicationRuntime(ROOT, Path(self.directory.name))
        self.classification = patch.object(self.runtime, "classify", return_value={"status": "classified", "predicted_class": "acid_reducer", "model": "fixture"})
        self.classification.start()
        self.client = TestClient(create_app(self.runtime))

    def tearDown(self):
        self.client.close()
        self.classification.stop()
        self.directory.cleanup()

    def test_health_model_list_and_openapi_are_available(self):
        self.assertEqual(self.client.get("/health").status_code, 200)
        self.assertEqual(len(self.client.get("/models").json()["models"]), 5)
        self.assertIn("/ask", self.client.get("/openapi.json").json()["paths"])

    def test_document_upload_extract_and_clear_workflow(self):
        response = self.client.post("/documents/upload", files={"file": ("source.txt", TEXT.encode(), "text/plain")})
        self.assertEqual(response.status_code, 200)
        record = response.json()
        self.assertTrue(record["extraction"]["fields"]["dosage"]["present"])
        identifier = record["document_id"]
        self.assertEqual(self.client.get("/documents/" + identifier).json()["text"], record["text"])
        self.assertEqual(self.client.delete("/documents/" + identifier).status_code, 200)
        self.assertEqual(self.client.get("/documents/" + identifier).status_code, 404)

    def test_repeated_identical_text_has_distinct_opaque_document_handles(self):
        first = self.client.post("/documents", json={"text": TEXT}).json()
        second = self.client.post("/documents", json={"text": TEXT}).json()
        self.assertNotEqual(first["document_id"], second["document_id"])
        self.assertEqual(first["source_hash"], second["source_hash"])
        self.client.delete("/documents/" + first["document_id"])
        self.assertEqual(self.client.get("/documents/" + second["document_id"]).status_code, 200)

    def test_ask_unknown_or_deleted_document_does_not_use_another_document(self):
        self.client.post("/documents", json={"text": TEXT})
        result = self.client.post("/ask", json={"document_id": "missing", "question": "Where is the amber widget?"})
        self.assertEqual(result.status_code, 404)

    def test_expired_document_is_released(self):
        record = self.client.post("/documents", json={"text": TEXT}).json()
        self.runtime.documents[record["document_id"]]["created_monotonic"] = time.monotonic() - SETTINGS["document_ttl_seconds"] - 1
        self.assertEqual(self.client.get("/documents/" + record["document_id"]).status_code, 404)

    def test_invalid_empty_input_is_rejected(self):
        self.assertEqual(self.client.post("/documents", json={"text": ""}).status_code, 422)
        self.assertEqual(self.client.post("/extract", json={"text": " "}).status_code, 422)

    def test_source_only_extract_endpoint_has_no_classifier_dependency(self):
        response = self.client.post("/extract", json={"text": TEXT})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["method"], "source-heading section extraction")

    def test_large_docx_upload_and_simple_questions_return_concise_grounded_answers(self):
        config = json.loads((ROOT / "configs" / "phase4.json").read_text(encoding="utf-8"))
        self.runtime.assistant = DocumentAssistant(ROOT, config, embedder=FakeEmbedder(Path(self.directory.name)),
            answerability=FakeAnswerability(), generator=FakeGenerator("unsupported answer"), storage_root=Path(self.directory.name))
        document = Document()
        fixture = (ROOT / "tests" / "fixtures" / "qa" / "long_patient_leaflet.txt").read_text(encoding="utf-8")
        for line in fixture.splitlines():
            if line.strip():
                document.add_paragraph(line)
        buffer = io.BytesIO()
        document.save(buffer)
        upload = self.client.post("/documents/upload", files={"file": ("large_test.docx", buffer.getvalue())})
        self.assertEqual(upload.status_code, 200)
        record = upload.json()
        for question, expected in (("can I take the medicine if I'm pregnant?", "only if the potential benefit"),
                                   ("dosage", "one 5 mg tablet"), ("how often should I take it?", "twice daily"),
                                   ("how long should I take it?", "7 days only"), ("price", "$24.95")):
            with self.subTest(question=question):
                response = self.client.post("/ask", json={"document_id": record["document_id"], "question": question})
                self.assertEqual(response.status_code, 200)
                result = response.json()
                self.assertEqual(result["status"], "answered")
                self.assertIn(expected, result["answer"])
                self.assertLess(len(result["answer"]), 600)
                self.assertTrue(response_supported(result, record["text"]))

    def test_qa_document_switch_does_not_copy_price_or_pregnancy_condition(self):
        config = json.loads((ROOT / "configs" / "phase4.json").read_text(encoding="utf-8"))
        self.runtime.assistant = DocumentAssistant(ROOT, config, embedder=FakeEmbedder(Path(self.directory.name)),
            answerability=FakeAnswerability(), generator=FakeGenerator("unsupported answer"), storage_root=Path(self.directory.name))
        records = []
        for filename in ("long_patient_leaflet.txt", "long_uncertain_leaflet.txt"):
            text = (ROOT / "tests" / "fixtures" / "qa" / filename).read_text(encoding="utf-8")
            records.append(self.client.post("/documents", json={"text": text, "source_name": filename}).json())
        first = self.client.post("/ask", json={"document_id": records[0]["document_id"], "question": "price"}).json()
        second = self.client.post("/ask", json={"document_id": records[1]["document_id"], "question": "price"}).json()
        self.assertIn("$24.95", first["answer"])
        self.assertEqual(second["status"], "not_found")
        response = self.client.post("/ask", json={"document_id": records[1]["document_id"], "question": "Can I take it if I'm pregnant?"}).json()
        self.assertEqual(response["answer_polarity"], "uncertain")
        self.assertNotIn("potential benefit", response["answer"])
        self.assertTrue(response_supported(response, records[1]["text"]))


if __name__ == "__main__":
    unittest.main()
