"""Create larger, clearly fictional upload fixtures without changing real example labels."""

import json
from pathlib import Path

from docx import Document

from medical_doc_assistant.cleaning import clean_text
from medical_doc_assistant.ingestion import ALIASES, read_document_file
from medical_doc_assistant.storage import file_sha256, write_json


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    source = root / "tests" / "fixtures" / "qa"
    target = root / "upload_examples" / "qa_tests"
    target.mkdir(exist_ok=True)
    settings = json.loads((root / "configs" / "phase5.json").read_text(encoding="utf-8"))
    manifest = []
    for path in sorted(source.glob("*.txt")):
        text = path.read_text(encoding="utf-8")
        canonical = clean_text(text)
        copy = target / path.name
        copy.write_text(text, encoding="utf-8")
        document = Document()
        for line in text.splitlines():
            if not line.strip():
                continue
            if line.strip().lower() in ALIASES:
                document.add_heading(line, level=1)
            else:
                document.add_paragraph(line)
        docx = target / (path.stem + ".docx")
        document.save(docx)
        if read_document_file(docx.name, docx.read_bytes(), settings) != canonical:
            raise AssertionError("Generated DOCX extraction changed the fixture: " + docx.name)
        manifest.append({"source_fixture": str(path.relative_to(root)).replace("\\", "/"),
                         "txt": str(copy.relative_to(root)).replace("\\", "/"),
                         "docx": str(docx.relative_to(root)).replace("\\", "/"),
                         "words": len(canonical.split()), "characters": len(canonical),
                         "docx_sha256": file_sha256(docx), "is_fictional_test_data": True,
                         "used_for_training": False})
    write_json(target / "manifest.json", {"purpose": "Large fictional QA upload tests; excluded from training", "documents": manifest})
    print(json.dumps({"generated": len(manifest), "directory": str(target),
                      "documents": [{"name": item["docx"], "words": item["words"], "characters": item["characters"]} for item in manifest]}, indent=2))


if __name__ == "__main__":
    main()
