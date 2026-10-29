"""Create upload-ready PDF/DOCX/TXT files from the project's real FDA labels."""

import html
import json
import re
import zipfile
from pathlib import Path

import matplotlib
import pandas as pd
from docx import Document
from docx.shared import Pt
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from medical_doc_assistant.ingestion import read_document_file
from medical_doc_assistant.storage import file_sha256, write_json

CHOICES = [
    ("01_Esomeprazole", "ESOMEPRAZOLE MAGNESIUM DIHYDRATE"),
    ("02_Aspirin", "ASPIRIN"),
    ("03_Nystatin", "NYSTATIN"),
    ("04_Erythromycin", "ERYTHROMYCIN"),
    ("05_Desoximetasone", "DESOXIMETASONE"),
    ("06_Pindolol", "PINDOLOL"),
]


def normalize_for_comparison(text: str) -> str:
    # PDF layout may wrap a long source word across lines; compare all source
    # characters while ignoring layout whitespace, without changing the files.
    return re.sub(r"\s+", "", text)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    folder = root / "upload_examples"
    for format_name in ("pdf", "docx", "txt"):
        (folder / format_name).mkdir(parents=True, exist_ok=True)
    cases = json.loads((root / "data" / "evaluation" / "phase4_qa_cases.json").read_text(encoding="utf-8"))
    dataset_path = root / "data" / "processed" / "medicine_labels.parquet"
    frame = pd.read_parquet(dataset_path).set_index("record_id")
    settings = json.loads((root / "configs" / "phase5.json").read_text(encoding="utf-8"))
    fonts = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    pdfmetrics.registerFont(TTFont("LabelSans", str(fonts / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("LabelSansBold", str(fonts / "DejaVuSans-Bold.ttf")))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="LabelBody", fontName="LabelSans", fontSize=9, leading=13, spaceAfter=5))
    styles.add(ParagraphStyle(name="LabelHeading", fontName="LabelSansBold", fontSize=12, leading=16, spaceBefore=10, spaceAfter=6))
    samples = []
    for filename, generic in CHOICES:
        source = next(source for source in cases["sources"] if generic in json.loads(source["generic_names"]))
        row = frame.loc[source["record_id"]]
        text = row["document_text"]
        paths = {kind: folder / kind / (filename + "." + kind) for kind in ("txt", "docx", "pdf")}
        paths["txt"].write_text(text, encoding="utf-8")
        docx = Document()
        docx.styles["Normal"].font.name = "Calibri"
        docx.styles["Normal"].font.size = Pt(10)
        elements = []
        for line in text.splitlines():
            if not line.strip():
                docx.add_paragraph()
                elements.append(Spacer(1, 5))
                continue
            is_heading = len(line) < 85 and line.replace(" ", "").isalpha() and line == line.title()
            if is_heading:
                docx.add_heading(line, level=1)
            else:
                docx.add_paragraph(line)
            elements.append(Paragraph(html.escape(line), styles["LabelHeading"] if is_heading else styles["LabelBody"]))
        docx.save(paths["docx"])
        SimpleDocTemplate(str(paths["pdf"]), title=filename.replace("_", " "), author="Source: openFDA", topMargin=35, bottomMargin=35).build(elements)
        for kind, path in paths.items():
            parsed = read_document_file(path.name, path.read_bytes(), settings)
            if normalize_for_comparison(parsed) != normalize_for_comparison(text):
                raise ValueError(f"{path.name} does not preserve the source text after ingestion.")
        samples.append({"name": filename.split("_", 1)[1], "source_generic_name": generic, "source_record_id": source["record_id"],
                        "source_category": row["class_label"], "source_url": row["source_url"], "is_synthetic": False,
                        "characters": len(text), "words": int(row["word_count"]),
                        "files": {kind: str(path.relative_to(folder)).replace("\\", "/") for kind, path in paths.items()},
                        "sha256": {kind: file_sha256(path) for kind, path in paths.items()}})
    write_json(folder / "manifest.json", {"source": "Real openFDA medicine labels from the frozen project corpus", "dataset_sha256": file_sha256(dataset_path),
                                         "medical_text_is_synthetic": False, "formats": ["pdf", "docx", "txt"], "samples": samples})
    table = "\n".join(f"| {sample['name']} | `{sample['files']['pdf']}` | `{sample['source_category']}` |" for sample in samples)
    guide = f"""# Files to upload to the Medical Document Assistant

Start with **`pdf/01_Esomeprazole.pdf`**. It is a real medicine-label example and
works with the application's upload, section extraction, classification, and chat.

## How to use

1. Open the app and choose **Upload file**.
2. Select one PDF from this folder's `pdf/` subfolder.
3. Click **Analyze document**.
4. Review the category and source sections, then ask a question about that document.

The `docx/` and `txt/` versions contain the same source text. You can upload them,
or open a TXT file and paste its contents into **Paste text**.

## Included real medicine labels

| Medicine | Suggested upload | Source-derived category |
|---|---|---|
{table}

## Questions to try

For **Esomeprazole**:

- How long can this product take for full effect?
- What does the document say about dosage and administration?
- What does the document say about pregnancy or breast-feeding?
- What is the retail price? (Tests the missing-answer fallback.)

For the other files:

- What does the document say about warnings?
- What does the document say about contraindications?
- What does the document say about adverse reactions and side effects?
- According to the dosage directions, how often is this product used?

Some files lack a dedicated section. Missing source text stays empty, and a model
may decline a question when it cannot establish sufficient evidence.

## Provenance

These are upload-ready renderings of the project's **real cleaned openFDA labels**,
with the same medical text in every format. They are not official DailyMed PDF
layouts. Source IDs/URLs and file checksums are in `manifest.json`; no patient
names, invented prescriptions, or fabricated medical instructions were added.

Data provided by the U.S. Food and Drug Administration (<https://open.fda.gov>).
Use source terms: <https://open.fda.gov/terms/>. This project is informational only
and is not medical advice.

## Training data

The full model-training corpus is `../data/processed/medicine_labels.parquet` or
`../data/processed/medicine_labels.csv.gz`. For app testing, choose the individual
leaflet files above. All 18 files were parsed with the application's ingestion
code and checked against the source text.
"""
    (folder / "README.md").write_text(guide, encoding="utf-8")
    package = root / "deliverables" / "Upload_Examples.zip"
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in folder.rglob("*"):
            if path.is_file():
                archive.write(path, Path("upload_examples") / path.relative_to(folder))
    print(json.dumps({"folder": str(folder), "zip": str(package), "documents": len(samples), "upload_files": len(samples) * 3,
                      "all_source_text_round_trips_verified": True}, indent=2))


if __name__ == "__main__":
    main()
