"""Generate architecture visuals, final slide deck/PDF, speaker notes, and source handoff."""

import json
import textwrap
import zipfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.pdfgen import canvas

from medical_doc_assistant.storage import write_json


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    destination = root / "deliverables"
    destination.mkdir(exist_ok=True)
    comparison = pd.read_csv(root / "reports" / "phase4" / "bert" / "model_comparison.csv")
    integration = json.loads((root / "reports" / "phase5" / "integration_results.json").read_text(encoding="utf-8"))
    live_file = root / "reports" / "phase5" / "live_demo.json"
    live = json.loads(live_file.read_text(encoding="utf-8")) if live_file.exists() else {}

    fig, ax = plt.subplots(figsize=(14, 7))
    ax.set(xlim=(0, 14), ylim=(0, 7))
    ax.axis("off")
    boxes = [
        (0.3, 2.9, "Upload / paste", "PDF · DOCX · text"),
        (3.0, 2.9, "Streamlit", "Source results + chat"),
        (5.7, 2.9, "FastAPI", "Opaque document handles"),
        (8.4, 4.9, "TextCNN / models", "Validation-selected classifier"),
        (8.4, 2.9, "Source extraction", "Exact sections + offsets"),
        (8.4, 0.9, "Per-document FAISS", "MiniLM semantic retrieval"),
        (11.1, 0.9, "Checked QA", "Answerability + FLAN-T5"),
        (11.1, 2.9, "Evidence / fallback", "Source quote · no invented facts"),
    ]
    for x, y, title, subtitle in boxes:
        ax.add_patch(FancyBboxPatch((x, y), 2.35, 1.1, boxstyle="round,pad=0.05", facecolor="#eef1f8", edgecolor="#5e5aa1", linewidth=1.5))
        ax.text(x + 1.175, y + .72, title, ha="center", va="center", fontsize=11, weight="bold", color="#263450")
        ax.text(x + 1.175, y + .3, subtitle, ha="center", va="center", fontsize=8, color="#526078")
    for first, second in [(0, 1), (1, 2), (2, 3), (2, 4), (2, 5), (5, 6), (6, 7)]:
        x1, y1, _, _ = boxes[first]
        x2, y2, _, _ = boxes[second]
        ax.annotate("", xy=(x2, y2 + .55), xytext=(x1 + 2.35, y1 + .55), arrowprops={"arrowstyle": "->", "color": "#687995", "lw": 1.4})
    ax.text(7, 6.6, "Medical Document Assistant — serving architecture", ha="center", fontsize=19, weight="bold", color="#263450")
    fig.savefig(root / "docs" / "architecture.png", dpi=160, bbox_inches="tight")
    fig.savefig(root / "docs" / "architecture.svg", bbox_inches="tight")
    plt.close(fig)

    slides = [
        ("Medical Document Assistant", ["AI Engineering capstone · Python → ML → Deep Learning → NLP → Web app", "English medicine-label classification and source-quoted document QA", "Informational only — not medical advice"],
         "Introduce the project as an educational document assistant. A user uploads or pastes text, sees an automatically predicted medicine category and source sections, then asks explicit questions with evidence. The main engineering objective is an end-to-end system: real data, reproducible model comparisons, retrieval, verified answers, and a usable interface. Emphasize that the app quotes a supplied source rather than recommending treatment. The evaluated language is English; Arabic and image OCR are scoped extensions. This presentation follows the five course phases and finishes with a live example."),
        ("Architecture and application flow", ["Streamlit frontend → FastAPI backend", "Source-heading extraction and cached classifiers", "Independent FAISS index → answerability → local generation → source verification"],
         "Walk through the diagram from the user's file to its canonical text. FastAPI assigns an opaque handle, so the frontend keeps its questions attached to the current document. Classification and extraction are separate from QA: a wrong category does not change which source document is searched. Retrieval uses a separate FAISS index for each text. A local instruction model proposes an answer span, but a verifier returns only a supported quotation or the fallback. CPU models are cached, and shared tokenizer inference is serialized. The Docker container supervises the two web processes."),
        ("Real data and fair evaluation", ["3,572 raw openFDA records → 2,471 cleaned English labels", "Six explicit EPC-derived course categories", "Frozen split: 1,762 train / 367 validation / 342 test", "100 ingredient-aware components; source and split hashes recorded"],
         "Explain that no medical training documents were invented. We downloaded real structured labels from openFDA, preserved the clinical sections, normalized HTML and encoding, and removed identical documents and ambiguous target combinations. Labels are automatic pharmacologic-class metadata targets, not expert-confirmed diagnoses or indications. Repackagers often supply similar text, so random row splitting would overestimate performance. The project freezes ingredient-aware components, with exact RxNorm links and duplicate checks, before training. Every model uses the same record assignment. The source manifests and hashes make the comparison auditable and reproducible."),
        ("Classifier comparison", [f"{row.display_name}: accuracy {row.test_accuracy:.2%} · macro F1 {row.test_macro_f1:.4f}" for row in comparison.itertuples()],
         "Describe the three TF-IDF baselines, the learned-embedding TextCNN, and the compact BERT. Hyperparameters use training and validation, with test results used for reporting. TextCNN was the validation-selected default and obtained the strongest recorded test result. The compact four-layer BERT did not beat it. Capacity and input budgets differ: TF-IDF reads full text, the CNN samples up to 2,048 regex tokens, and BERT samples 508 content WordPieces. All measured outcomes are retained. This is a comparison on a small source-derived snapshot, not a universal medical accuracy claim."),
        ("Grounded QA and honest testing", ["MiniLM embeddings · per-document FAISS · local FLAN-T5", "Answers are source quotations with exact character offsets", "Missing evidence: 'I cannot find that in the document.'", "20 regression cases / 3 evaluation leaflets: 11 answerable, 9 missing"],
         "Explain the answerability and grounding checks before showing a response. Broad section FAQs can retrieve explicit source headings; specific queries use semantic retrieval and a SQuAD2 model. Generated text is accepted only as a checked span, with source clause context included. The initial QA evaluation exposed an important bug: a frequency question received a strength description. A generic temporal-answer check corrected that, and both the initial and updated results are preserved. The updated small template-based benchmark passed its source-evidence and fallback checks. It is a regression check, not clinical validation or a fresh untouched medical QA estimate."),
        ("Live document demonstration", ["Load the real FDA example or upload a text-based PDF", "Inspect the predicted category and dosage/warning source sections", "Ask: 'How long can this product take for full effect?'", "Ask an absent fact, then clear the document"],
         "Use about one minute for the live demonstration. Load the built-in real leaflet and point to the predicted category, document word count, and exact dosage section. Ask the full-effect question and open its source-evidence expander to show the quote and character positions. Then ask for retail price: this fact is absent, and the app should return the missing-answer phrase. Finally clear or replace the document to show conversation isolation. If network access is unavailable, use the captured browser screenshots in the report folder. Avoid interpreting the leaflet as individual medical advice; this demo shows source handling."),
        ("Verification and deployment", [f"Real API/browser integration: {len(integration['checks'])} checks passed", "Upload, classification, extraction, citations, fallbacks, and clearing tested", "Docker image + Compose + Hugging Face Docker Space package", "Temporary review URL and persistent-hosting handoff documented"],
         "Summarize verification at two levels. Software tests cover data isolation, padding, serialization, typed answer checks, ingestion, and API document handling. Real integration exercises the trained models, PDF and DOCX uploads, and browser chat flows. Docker packaging includes CPU PyTorch and immutable public model downloads. A deployment script stages the app for Hugging Face Spaces using the owner's local login; credentials are not placed in source. A temporary Cloudflare link can expose the running local app for review, but depends on that machine and tunnel remaining active. Persistent hosting is a separate account-backed publication step."),
        ("Conclusion and next improvements", ["Complete educational data-to-application pipeline", "Strongest recorded classifier: TextCNN · 98.54% accuracy / 0.9876 macro F1", "Explicit limits: source-derived labels, small QA benchmark, English text", "Extensions: evaluated Arabic models, OCR, broader QA cases, persistent hosting"],
         "Close with the engineering lessons. Strong TF-IDF baselines matter; a larger pretrained model is not automatically better under a compact runtime and short context. Related-document splitting and source hashes improve the validity of a model comparison. Source support alone is not enough for QA: the answer must also match the requested type, and negative questions need evaluation. The final application is usable and modular, but the source labels and clinical claims remain unadjudicated. Future work should broaden independent QA references, add tested Arabic and OCR support, and operate a persistent deployment with appropriate identity and storage design. Invite questions about the architecture or evaluation."),
    ]
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for index, (title, bullets, notes) in enumerate(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for position, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if position == 0 else frame.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = Pt(22 if index != 3 else 19)
            paragraph.space_after = Pt(14)
        slide.notes_slide.notes_text_frame.text = notes
        if index == 1:
            frame.text = "Upload → API → source sections / classifier → document-only retrieval → checked answer"
            frame.paragraphs[0].font.size = Pt(18)
            slide.shapes.add_picture(str(root / "docs" / "architecture.png"), Inches(1), Inches(1.8), height=Inches(5.2))
        if index == 5 and (root / "reports" / "phase5" / "screenshots" / "grounded_chat.png").exists():
            # The full browser screenshot is included as a separate review asset; the slide remains readable.
            slide.notes_slide.notes_text_frame.text += "\nScreenshot: reports/phase5/screenshots/grounded_chat.png"
    deck.save(destination / "Final_Presentation.pptx")
    pdf = canvas.Canvas(str(destination / "Final_Presentation.pdf"), pagesize=(960, 540))
    for index, (title, bullets, _) in enumerate(slides):
        pdf.setFillColor(colors.HexColor("#263450"))
        pdf.setFont("Helvetica-Bold", 26)
        pdf.drawString(45, 490, title)
        pdf.setFont("Helvetica", 17)
        y = 430
        for bullet in ([] if index == 1 else bullets):
            for line in textwrap.wrap(bullet, 85):
                pdf.drawString(48, y, line)
                y -= 25
            y -= 12
        if index == 1:
            pdf.drawImage(str(root / "docs" / "architecture.png"), 45, 55, width=870, height=400, preserveAspectRatio=True)
        pdf.setFont("Helvetica", 9)
        pdf.drawString(45, 15, "Medical Document Assistant | Informational only | " + str(index + 1))
        pdf.showPage()
    pdf.save()
    notes = "# Presentation speaker notes (approximately 6–7 minutes including demo)\n\n"
    for index, (title, _, note) in enumerate(slides, 1):
        notes += f"## Slide {index}: {title}\n\n{note}\n\n"
    (destination / "Speaker_Notes.md").write_text(notes, encoding="utf-8")
    with zipfile.ZipFile(destination / "Source_Code_Handoff.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
        directories = ["src", "configs", "scripts", "tests", "docs", "deployment"]
        for directory in directories:
            for path in (root / directory).rglob("*"):
                if path.is_file() and "__pycache__" not in path.parts and not path.name.endswith(".pyc"):
                    archive.write(path, path.relative_to(root))
        for filename in ["README.md", "pyproject.toml", ".gitignore", ".dockerignore", "Dockerfile", "docker-compose.yml", "streamlit_app.py", "Final_Project_Plan.pdf"]:
            archive.write(root / filename, filename)
        for path in root.glob("requirements-phase*.txt"):
            archive.write(path, path.name)
    metadata = {"slides": len(slides), "speaker_note_words": sum(len(note.split()) for _, _, note in slides),
                "live_demo": live.get("url"), "files": ["Final_Presentation.pptx", "Final_Presentation.pdf", "Speaker_Notes.md", "Source_Code_Handoff.zip"]}
    write_json(destination / "manifest.json", metadata)
    print(json.dumps(metadata, indent=2))


if __name__ == "__main__":
    main()
